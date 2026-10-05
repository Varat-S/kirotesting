"""XLSX / CSV financial-table parser (task 3.2).

Parses a tabular financial statement where the first column holds a row label
(line item) and subsequent columns hold period values. Each cell becomes a
:class:`~app.schemas.evidence.CanonicalFact` carrying:

* ``source_label`` / ``name`` = the row (line-item) label;
* a :class:`~app.schemas.evidence.SourceRef` locating the sheet (``table``),
  the ``row_label`` and the ``cell`` (e.g. ``B2``);
* ``scale`` and ``currency`` passed in by the caller (statement header), with
  raw vs normalized values preserved separately via
  :mod:`app.services.extraction.normalization`;
* the reporting ``period`` resolved from the column header.

Both XLSX (via :mod:`openpyxl`, if installed) and CSV (via the stdlib) are
supported through a common path that normalizes the input into a list of rows.

Error handling (Req 4.5, 23.2): a requested sheet or label column that does not
exist raises :class:`~app.services.extraction.base.ParserError`. An individual
empty cell is emitted as a MISSING fact (never zero / never fabricated).
"""

from __future__ import annotations

import csv
import io
import re
import uuid

from app.schemas.enums import ExtractionMethod, FactStatus
from app.schemas.evidence import CanonicalFact, SourceRef
from app.services.extraction.base import (
    BaseParser,
    ParserAuditHook,
    ParserError,
    ParseResult,
    ParserMetadata,
)
from app.services.extraction.normalization import (
    NormalizationError,
    normalize_amount,
    normalize_period,
)

_CSV_METADATA = ParserMetadata(
    name="csv_parser", version="1.0.0", extraction_method=ExtractionMethod.CSV
)
_XLSX_METADATA = ParserMetadata(
    name="xlsx_parser", version="1.0.0", extraction_method=ExtractionMethod.XLSX
)


def _column_letter(index_zero_based: int) -> str:
    """Return a spreadsheet column letter for a 0-based index (0 -> 'A')."""
    index = index_zero_based
    letters = ""
    index += 1
    while index > 0:
        index, rem = divmod(index - 1, 26)
        letters = chr(ord("A") + rem) + letters
    return letters


class TableParser(BaseParser):
    """Shared implementation for CSV and XLSX financial tables."""

    def __init__(self, metadata: ParserMetadata) -> None:
        self._metadata = metadata

    @property
    def metadata(self) -> ParserMetadata:
        return self._metadata

    def parse_rows(
        self,
        rows: list[list[str]],
        *,
        document_id: str,
        sheet_name: str,
        scale: str | None = None,
        currency: str | None = None,
        period_headers: dict[str, dict] | None = None,
        fiscal_year_end_month: int | None = None,
        target_scale: str = "millions",
        audit_hook: ParserAuditHook | None = None,
    ) -> ParseResult:
        """Parse an in-memory table (header row + label-first data rows).

        ``period_headers`` maps a column header string to a dict of
        :func:`~app.services.extraction.normalization.normalize_period` kwargs,
        letting the caller declare each period column deterministically.
        """
        if not rows:
            raise ParserError("Empty table: no header row present.")
        header = rows[0]
        if len(header) < 2:
            raise ParserError(
                "Table must have a label column plus at least one value column."
            )

        result = ParseResult(metadata=self.metadata)
        period_headers = period_headers or {}

        for r_index, row in enumerate(rows[1:], start=2):
            if not row or not (row[0] or "").strip():
                continue  # blank/spacer row
            row_label = row[0].strip()
            for c_index in range(1, len(header)):
                col_header = (header[c_index] or "").strip()
                cell_text = (row[c_index] if c_index < len(row) else "") or ""
                cell_text = cell_text.strip()
                cell_ref = f"{_column_letter(c_index)}{r_index}"
                source_ref = SourceRef(
                    document_id=document_id,
                    table=sheet_name,
                    row_label=row_label,
                    cell=cell_ref,
                )
                period = self._resolve_period(
                    col_header, period_headers, fiscal_year_end_month
                )

                if not cell_text:
                    # Empty cell -> explicit MISSING, never zero (Req 4.4).
                    result.facts.append(
                        CanonicalFact(
                            fact_id=str(uuid.uuid4()),
                            name=row_label,
                            source_label=row_label,
                            status=FactStatus.MISSING,
                            extraction_method=self.metadata.extraction_method,
                            period_start=period.period_start if period else None,
                            period_end=period.period_end if period else None,
                            period_type=period.period_type if period else None,
                            fiscal_year=period.fiscal_year if period else None,
                            source_refs=[source_ref],
                            created_by=self.metadata.created_by,
                        )
                    )
                    continue

                normalized = None
                try:
                    normalized = normalize_amount(
                        cell_text,
                        scale=scale,
                        currency=currency,
                        target_scale=target_scale,
                    )
                except NormalizationError:
                    normalized = None

                result.facts.append(
                    CanonicalFact(
                        fact_id=str(uuid.uuid4()),
                        name=row_label,
                        source_label=row_label,
                        raw_value=cell_text,
                        raw_unit=currency,
                        normalized_value=normalized.value if normalized else None,
                        normalized_unit=normalized.unit if normalized else None,
                        currency=currency,
                        scale=scale,
                        period_start=period.period_start if period else None,
                        period_end=period.period_end if period else None,
                        period_type=period.period_type if period else None,
                        fiscal_year=period.fiscal_year if period else None,
                        normalization_method=normalized.method if normalized else None,
                        status=FactStatus.UNVERIFIED,
                        extraction_method=self.metadata.extraction_method,
                        source_refs=[source_ref],
                        created_by=self.metadata.created_by,
                    )
                )

        self._log(result, audit_hook, document_id=document_id, sheet=sheet_name)
        return result

    def _resolve_period(
        self,
        col_header: str,
        period_headers: dict[str, dict],
        fiscal_year_end_month=None,
    ):
        spec = period_headers.get(col_header)
        if (
            spec is None
            and fiscal_year_end_month is not None
            and re.fullmatch(r"(?:FY\s*)?((?:19|20)\d{2})", col_header, re.I)
        ):
            spec = {
                "fiscal_year": col_header,
                "fiscal_year_end_month": fiscal_year_end_month,
            }
        if spec is None:
            return None
        try:
            return normalize_period(**spec)
        except NormalizationError:
            return None

    def parse(self, *args, **kwargs) -> ParseResult:  # pragma: no cover - delegated
        raise NotImplementedError("Use CsvParser.parse or XlsxParser.parse.")


class CsvParser(TableParser):
    """Parse a CSV financial table."""

    def __init__(self) -> None:
        super().__init__(_CSV_METADATA)

    def parse(
        self,
        data: bytes | str,
        *,
        document_id: str,
        sheet_name: str = "csv",
        scale: str | None = None,
        currency: str | None = None,
        period_headers: dict[str, dict] | None = None,
        fiscal_year_end_month: int | None = None,
        target_scale: str = "millions",
        audit_hook: ParserAuditHook | None = None,
    ) -> ParseResult:
        text = data.decode("utf-8") if isinstance(data, bytes) else data
        try:
            reader = csv.reader(io.StringIO(text))
            rows = [list(r) for r in reader]
        except csv.Error as exc:
            raise ParserError(f"Malformed CSV: {exc}") from exc
        return self.parse_rows(
            rows,
            document_id=document_id,
            sheet_name=sheet_name,
            scale=scale,
            currency=currency,
            period_headers=period_headers,
            fiscal_year_end_month=fiscal_year_end_month,
            target_scale=target_scale,
            audit_hook=audit_hook,
        )


class XlsxParser(TableParser):
    """Parse an XLSX financial table (requires :mod:`openpyxl`)."""

    def __init__(self) -> None:
        super().__init__(_XLSX_METADATA)

    def parse(
        self,
        data: bytes,
        *,
        document_id: str,
        sheet_name: str | None = None,
        scale: str | None = None,
        currency: str | None = None,
        period_headers: dict[str, dict] | None = None,
        fiscal_year_end_month: int | None = None,
        target_scale: str = "millions",
        audit_hook: ParserAuditHook | None = None,
    ) -> ParseResult:
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - env dependent
            raise ParserError(
                "openpyxl is required to parse XLSX files but is not installed."
            ) from exc

        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=True
            )
        except Exception as exc:  # openpyxl raises various errors
            raise ParserError(f"Unreadable XLSX workbook: {exc}") from exc

        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise ParserError(
                    f"Sheet {sheet_name!r} not found; available: {workbook.sheetnames}"
                )
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active
            sheet_name = worksheet.title

        rows: list[list[str]] = []
        for raw_row in worksheet.iter_rows(values_only=True):
            rows.append(["" if v is None else str(v) for v in raw_row])

        return self.parse_rows(
            rows,
            document_id=document_id,
            sheet_name=sheet_name,
            scale=scale,
            currency=currency,
            period_headers=period_headers,
            fiscal_year_end_month=fiscal_year_end_month,
            target_scale=target_scale,
            audit_hook=audit_hook,
        )
