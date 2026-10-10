"""Healthcare benchmark workbook importer + validator (deterministic).

Reads the ORIGINAL ``.xlsx`` bytes (never modifies them), records the SHA-256 of
the source file, and produces two machine-readable artifacts:

* a normalized, content-hashed **reference dataset** (industry aggregates,
  SIC table, company identity list, assumptions, source documentation); and
* a **validation report** listing every issue found and, for every imported
  benchmark value, whether it is eligible for its declared use or excluded and
  why.

Every value carries exactly one status:

``observed``          a recorded identity/classification fact (SIC rows, company list);
``source_aggregate``  an industry aggregate transcribed from the published source;
``derived``           computed by a workbook formula from source aggregates;
``assumed``           an editable assumption, or anything that depends on a
                      placeholder assumption;
``unavailable``       explicitly not provided by the workbook (``n/a`` / ``n/m``);
``invalid``           an error value, an unreconcilable cached result, an
                      unresolved external dependency, a blank mandatory cell, an
                      implausible value or an inconsistent formula.

Cached Excel results are not trusted: every formula is re-derived from the
workbook's literal inputs (:mod:`app.services.benchmarking.formula`) and compared
with the cached value. A reference into another workbook is rebound to the local
sheet of the same name ONLY when the external link's cached value equals the
local cell, and that correction is logged; otherwise the value is ``invalid``.

The workbook layout (sheet names, section markers, assumption cells, row ->
metric mapping) is DATA, supplied by the versioned sector configuration; no
benchmark value is hard-coded here.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from app.core.hashing import content_hash
from app.services.benchmarking.formula import (
    EXCEL_ERRORS,
    CellRef,
    ExcelError,
    UnsupportedFormula,
    evaluate,
    references,
    relative_form,
)

IMPORTER_VERSION = "healthcare-workbook-importer-1.0.0"
DATASET_SCHEMA_VERSION = "industry-reference-dataset-1.0"

VALUE_STATUSES = (
    "observed",
    "source_aggregate",
    "derived",
    "assumed",
    "unavailable",
    "invalid",
)

_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_MONTHS = (
    "january|february|march|april|may|june|july|august|september|october|"
    "november|december"
)
_VINTAGE = re.compile(rf"as of ({_MONTHS}) (\d{{4}})", re.IGNORECASE)
_UPDATED = re.compile(rf"last updated ({_MONTHS}) (\d{{1,2}}), (\d{{4}})", re.IGNORECASE)
_URL = re.compile(r"https?://\S+")


class WorkbookImportError(ValueError):
    """The file is not a readable workbook with the required structure."""


@dataclass
class Issue:
    severity: str  # error | warning | info
    code: str
    message: str
    sheet: str | None = None
    cell: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "severity": self.severity,
            "code": self.code,
            "sheet": self.sheet,
            "cell": self.cell,
            "message": self.message,
        }


@dataclass
class WorkbookImport:
    """Result of importing one workbook."""

    dataset: dict[str, Any]
    report: dict[str, Any]

    @property
    def dataset_hash(self) -> str:
        return self.dataset["dataset_hash"]

    @property
    def source_sha256(self) -> str:
        return self.dataset["source"]["sha256"]


@dataclass
class _External:
    """One external-workbook link and the values Excel cached for it."""

    index: int
    target: str | None
    sheet_names: list[str]
    cached: dict[tuple[str, str], Any] = field(default_factory=dict)
    refresh_errors: list[str] = field(default_factory=list)


class HealthcareWorkbookImporter:
    """Import + validate the healthcare benchmark workbook under a mapping."""

    def __init__(self, mapping: dict[str, Any], *, benchmark_metrics: dict[str, Any] | None = None) -> None:
        self._mapping = mapping
        self._metric_cfg = benchmark_metrics or {}
        self._issues: list[Issue] = []
        self._corrections: list[dict[str, Any]] = []

    # ------------------------------------------------------------------ API

    def import_bytes(self, data: bytes, *, filename: str = "workbook.xlsx") -> WorkbookImport:
        self._issues, self._corrections = [], []
        self._notes_column = None
        sha256 = hashlib.sha256(data).hexdigest()
        try:
            formulas = load_workbook(io.BytesIO(data), data_only=False)
            values = load_workbook(io.BytesIO(data), data_only=True)
            archive = zipfile.ZipFile(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001 - any failure is "not a workbook"
            raise WorkbookImportError(f"Unreadable workbook: {type(exc).__name__}.") from exc

        sheets = self._mapping["sheets"]
        missing = [name for name in sheets.values() if name not in formulas.sheetnames]
        if missing:
            raise WorkbookImportError(f"Workbook is missing required sheet(s): {missing}.")

        self._formulas, self._values = formulas, values
        self._externals = self._read_externals(archive)
        self._memo: dict[tuple[str, str], Any] = {}
        self._deps: dict[tuple[str, str], dict[str, Any]] = {}
        self._in_progress: set[tuple[str, str]] = set()
        self._assumptions = self._read_assumptions()
        self._scan_error_cells()

        industries = self._read_industries()
        metrics = self._read_benchmarks(industries)
        sic_codes = self._read_sic_codes()
        companies = self._read_companies(industries)
        documentation = self._read_documentation()
        vintage = self._read_vintage()
        self._cross_checks(metrics)
        self._check_sample_sizes(metrics, companies)

        dataset: dict[str, Any] = {
            "schema_version": DATASET_SCHEMA_VERSION,
            "importer_version": IMPORTER_VERSION,
            "mapping_hash": content_hash(self._mapping),
            "source": {
                "filename": filename,
                "sha256": sha256,
                "size_bytes": len(data),
                "sheets": list(formulas.sheetnames),
                "external_links": [
                    {
                        "index": e.index,
                        "target": e.target,
                        "sheet_names": e.sheet_names,
                        "refresh_errors": e.refresh_errors,
                    }
                    for e in self._externals.values()
                ],
            },
            "vintage": vintage,
            "industries": industries,
            "aggregate_nature": (
                "Cumulated industry aggregates (sum of numerators over sum of "
                "denominators). Not medians, percentiles or company-level records."
            ),
            "assumptions": self._assumptions,
            "documentation": documentation,
            "metrics": metrics,
            "sic_codes": sic_codes,
            "companies": companies,
            "correction_log": {
                "parent_sha256": sha256,
                "original_modified": False,
                "corrections": self._corrections,
            },
        }
        dataset["dataset_hash"] = content_hash(dataset)
        return WorkbookImport(dataset=dataset, report=self._report(dataset))

    # ------------------------------------------------------------ externals

    def _read_externals(self, archive: zipfile.ZipFile) -> dict[int, _External]:
        from lxml import etree

        out: dict[int, _External] = {}
        names = sorted(
            n for n in archive.namelist()
            if re.fullmatch(r"xl/externalLinks/externalLink\d+\.xml", n)
        )
        for name in names:
            index = int(re.search(r"(\d+)\.xml$", name).group(1))
            root = etree.fromstring(archive.read(name))
            sheet_names = [
                el.get("val") for el in root.iterfind(".//m:sheetNames/m:sheetName", _NS)
            ]
            target = None
            rel_name = f"xl/externalLinks/_rels/externalLink{index}.xml.rels"
            if rel_name in archive.namelist():
                rels = etree.fromstring(archive.read(rel_name))
                targets = [
                    r.get("Target") for r in rels.iterfind(f"{{{_REL_NS}}}Relationship")
                ]
                target = sorted(t for t in targets if t)[0] if targets else None
            ext = _External(index=index, target=target, sheet_names=sheet_names)
            for sheet_data in root.iterfind(".//m:sheetDataSet/m:sheetData", _NS):
                sheet_id = int(sheet_data.get("sheetId", "0"))
                sheet = sheet_names[sheet_id] if sheet_id < len(sheet_names) else None
                if sheet is None:
                    continue
                if sheet_data.get("refreshError") == "1":
                    ext.refresh_errors.append(sheet)
                for cell in sheet_data.iterfind(".//m:cell", _NS):
                    value = cell.find("m:v", _NS)
                    if value is None or value.text is None:
                        continue
                    try:
                        parsed: Any = float(value.text)
                    except ValueError:
                        parsed = value.text
                    ext.cached[(sheet, cell.get("r"))] = parsed
            out[index] = ext
            self._issue(
                "warning", "external_workbook_link",
                f"Workbook links to external file {target!r} "
                f"(sheets {sheet_names}); dependent formulas are re-derived locally.",
            )
        return out

    # ----------------------------------------------------------- evaluation

    def _raw(self, sheet: str, cell: str) -> Any:
        return self._formulas[sheet][cell].value

    def _cached(self, sheet: str, cell: str) -> Any:
        return self._values[sheet][cell].value

    def _resolve_external(self, ref: CellRef, origin: tuple[str, str]) -> Any:
        ext = self._externals.get(ref.external or -1)
        deps = self._deps[origin]
        target = ext.target if ext else None
        label = f"[{ref.external}]{ref.sheet}!{ref.cell}"
        deps["external"].add(label)
        policy = self._mapping.get("external_reference_policy")
        local_sheet = ref.sheet if ref.sheet in self._formulas.sheetnames else None
        if ext is None or local_sheet is None or policy != "rebind_to_local_sheet_when_cached_value_matches":
            deps["unresolved_external"].add(label)
            return ExcelError("#REF!")
        cached_external = ext.cached.get((ref.sheet, ref.cell))
        local = self._evaluate_cell(local_sheet, ref.cell)
        if cached_external is None or not _same(cached_external, local, 0.0, 0.0):
            deps["unresolved_external"].add(label)
            return ExcelError("#REF!")
        correction = {
            "kind": "external_reference_rebound_to_local",
            "reference": label,
            "external_target": target,
            "external_cached_value": cached_external,
            "local_reference": f"{local_sheet}!{ref.cell}",
            "local_value": local,
            "rationale": (
                "The external workbook is not available. The link's cached value "
                "equals the local cell on the sheet of the same name, so the "
                "reference is re-pointed to the local, documented input."
            ),
        }
        if correction not in self._corrections:
            self._corrections.append(correction)
        # Inherit the local cell's own dependencies (e.g. an assumption cell).
        self._merge_deps(origin, (local_sheet, ref.cell))
        deps["cells"].add(f"{local_sheet}!{ref.cell}")
        return local

    def _merge_deps(self, origin: tuple[str, str], child: tuple[str, str]) -> None:
        child_deps = self._deps.get(child)
        if child_deps is None:
            return
        for key in ("cells", "assumptions", "external", "unresolved_external"):
            self._deps[origin][key] |= child_deps[key]

    def _evaluate_cell(self, sheet: str, cell: str) -> Any:
        key = (sheet, cell)
        if key in self._memo:
            return self._memo[key]
        self._deps.setdefault(
            key,
            {"cells": set(), "assumptions": set(), "external": set(),
             "unresolved_external": set(), "unsupported": None},
        )
        for assumption in self._assumptions_index().get(key, []):
            self._deps[key]["assumptions"].add(assumption)
        raw = self._raw(sheet, cell)
        if not (isinstance(raw, str) and raw.startswith("=")):
            if isinstance(raw, str) and raw in EXCEL_ERRORS:
                raw = ExcelError(raw)
            self._memo[key] = raw
            return raw
        if key in self._in_progress:
            return ExcelError("#REF!")  # circular reference
        self._in_progress.add(key)

        def resolve(ref: CellRef) -> Any:
            if ref.external is not None:
                return self._resolve_external(ref, key)
            target = ref.key(sheet)
            if target[0] not in self._formulas.sheetnames:
                return ExcelError("#REF!")
            value = self._evaluate_cell(*target)
            self._deps[key]["cells"].add(f"{target[0]}!{target[1]}")
            self._merge_deps(key, target)
            return value

        try:
            result = evaluate(raw, resolve)
        except UnsupportedFormula as exc:
            self._deps[key]["unsupported"] = str(exc)
            result = ExcelError("#VALUE!")
        finally:
            self._in_progress.discard(key)
        self._memo[key] = result
        return result

    def _assumptions_index(self) -> dict[tuple[str, str], list[str]]:
        sheet = self._mapping["sheets"]["assumptions"]
        index: dict[tuple[str, str], list[str]] = {}
        for item in self._mapping["assumptions_sheet"]["inputs"]:
            index.setdefault((sheet, item["cell"]), []).append(item["key"])
        return index

    # ------------------------------------------------------------- sections

    def _read_assumptions(self) -> list[dict[str, Any]]:
        sheet = self._mapping["sheets"]["assumptions"]
        ws = self._formulas[sheet]
        out = []
        for item in self._mapping["assumptions_sheet"]["inputs"]:
            cell = item["cell"]
            row = ws[cell].row
            value = ws[cell].value
            if value is None:
                self._issue("error", "blank_mandatory_cell",
                            f"Assumption {item['key']!r} is blank.", sheet, cell)
            out.append(
                {
                    "key": item["key"],
                    "sheet": sheet,
                    "cell": cell,
                    "label": ws[f"A{row}"].value,
                    "value": value,
                    "meaning": ws[f"C{row}"].value,
                    "kind": item["kind"],
                    "status": "assumed",
                    "note": item.get("note"),
                }
            )
        return out

    def _scan_error_cells(self) -> None:
        """Flag every cell whose cached result or literal is an Excel error."""
        for ws in self._values.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value in EXCEL_ERRORS:
                        self._issue(
                            "error", "formula_error",
                            f"Cell holds Excel error {cell.value}.", ws.title, cell.coordinate,
                        )

    def _read_industries(self) -> list[str]:
        cfg = self._mapping["benchmarks_sheet"]
        ws = self._formulas[self._mapping["sheets"]["benchmarks"]]
        header = cfg["header_row"]
        industries: list[str] = []
        self._industry_columns: dict[str, str] = {}
        for cell in ws[header]:
            value = cell.value
            if cell.column_letter in {cfg["label_column"], cfg["definition_column"]}:
                continue
            if value is None or value == cfg["notes_header"]:
                if value == cfg["notes_header"]:
                    self._notes_column = cell.column_letter
                continue
            industries.append(str(value))
            self._industry_columns[str(value)] = cell.column_letter
        if self._notes_column is None:
            self._notes_column = get_column_letter(
                column_index_from_string(max(self._industry_columns.values())) + 1
            )
        expected = self._mapping["expected_industries"]
        if industries != expected:
            self._issue(
                "error", "industry_categories_mismatch",
                f"Industry columns {industries} differ from the mapped {expected}.",
                ws.title, f"A{header}",
            )
        return industries

    def _read_benchmarks(self, industries: list[str]) -> list[dict[str, Any]]:
        cfg = self._mapping["benchmarks_sheet"]
        sheet = self._mapping["sheets"]["benchmarks"]
        ws = self._formulas[sheet]
        tol = self._mapping["reconciliation"]
        row_map = {
            (r["section"], r["label"]): r["benchmark_metric_id"]
            for r in self._mapping["metric_rows"]
        }
        prefixes = cfg["section_prefixes"]
        section = None
        metrics: list[dict[str, Any]] = []
        seen_rows: set[tuple[str, str]] = set()
        for row in range(cfg["header_row"] + 1, ws.max_row + 1):
            label = ws[f"{cfg['label_column']}{row}"].value
            if label is None:
                continue
            label = str(label)
            matched = [s for s, p in prefixes.items() if label.startswith(p)]
            if matched:
                section = matched[0]
                continue
            metric_id = row_map.get((section, label))
            seen_rows.add((section, label))
            entry: dict[str, Any] = {
                "benchmark_metric_id": metric_id,
                "row": row,
                "section": section,
                "label": label,
                "definition": ws[f"{cfg['definition_column']}{row}"].value,
                "note": ws[f"{self._notes_column}{row}"].value,
                "sheet": sheet,
                "values": {},
            }
            cfg_metric = self._metric_cfg.get(metric_id or "", {})
            entry["declared_use"] = cfg_metric.get("use", "unmapped")
            entry["units"] = cfg_metric.get("units")
            forms: dict[str, str] = {}
            for industry in industries:
                column = self._industry_columns[industry]
                cell = f"{column}{row}"
                entry["values"][industry] = self._benchmark_value(
                    sheet, cell, section, cfg, tol, cfg_metric
                )
                raw = ws[cell].value
                if isinstance(raw, str) and raw.startswith("="):
                    forms[industry] = relative_form(raw, column)
            self._check_row_consistency(entry, forms, sheet)
            metrics.append(entry)
        for (sec, label), metric_id in sorted(row_map.items()):
            if (sec, label) not in seen_rows:
                self._issue(
                    "error", "mapped_row_missing",
                    f"Mapped benchmark row {label!r} (section {sec}, metric "
                    f"{metric_id!r}) is absent from the workbook.", sheet,
                )
        return metrics

    def _benchmark_value(self, sheet, cell, section, cfg, tol, cfg_metric) -> dict[str, Any]:
        raw = self._raw(sheet, cell)
        cached = self._cached(sheet, cell)
        is_formula = isinstance(raw, str) and raw.startswith("=")
        out: dict[str, Any] = {
            "sheet": sheet,
            "cell": cell,
            "formula": raw if is_formula else None,
            "cached_value": _jsonable(cached),
            "recalculated_value": None,
            "value": None,
            "status": None,
            "reasons": [],
            "dependencies": {"cells": [], "assumptions": [], "external": []},
            "reconciliation": None,
        }
        reasons: list[str] = out["reasons"]

        if raw is None:
            out["status"] = "invalid"
            reasons.append("blank_mandatory_cell")
            self._issue("error", "blank_mandatory_cell",
                        "Benchmark cell is blank.", sheet, cell)
            return self._finish(out, cfg_metric)

        if not is_formula:
            if isinstance(raw, bool) or not isinstance(raw, (int, float)):
                text = str(raw)
                if text in EXCEL_ERRORS:
                    out["status"] = "invalid"
                    reasons.append(f"formula_error:{text}")
                elif _is_unavailable(text, cfg):
                    out["status"] = "unavailable"
                    reasons.append("not_provided_by_workbook")
                else:
                    out["status"] = "invalid"
                    reasons.append("inappropriate_value:non_numeric_literal")
                    self._issue("error", "inappropriate_value",
                                f"Non-numeric benchmark literal {text!r}.", sheet, cell)
                return self._finish(out, cfg_metric)
            out["value"] = float(raw)
            if section == cfg["source_aggregate_section"]:
                out["status"] = "source_aggregate"
            else:
                # A hard-coded number where the layout expects a formula.
                out["status"] = "invalid"
                reasons.append("hard_coded_value_outside_source_section")
                self._issue("error", "hard_coded_value",
                            "Numeric literal outside the source-parameter section.",
                            sheet, cell)
            return self._finish(out, cfg_metric)

        recalculated = self._evaluate_cell(sheet, cell)
        deps = self._deps[(sheet, cell)]
        out["dependencies"] = {
            "cells": sorted(deps["cells"]),
            "assumptions": sorted(deps["assumptions"]),
            "external": sorted(deps["external"]),
        }
        out["recalculated_value"] = _jsonable(recalculated)
        cached_is_error = isinstance(cached, str) and cached in EXCEL_ERRORS

        if deps["unsupported"]:
            out["status"] = "invalid"
            reasons.append(f"unsupported_formula:{deps['unsupported']}")
            self._issue("error", "unsupported_formula",
                        f"Formula cannot be independently re-derived: {deps['unsupported']}",
                        sheet, cell)
            return self._finish(out, cfg_metric)
        if deps["unresolved_external"]:
            out["status"] = "invalid"
            reasons.append("unresolved_external_dependency:" + ",".join(sorted(deps["unresolved_external"])))
            self._issue("error", "unresolved_external_dependency",
                        "Formula depends on an external workbook that cannot be "
                        "reconciled to a local documented input.", sheet, cell)
            return self._finish(out, cfg_metric)
        if deps["external"]:
            reasons.append("external_reference_rebound_to_local")
            self._issue("warning", "external_dependency_rebound",
                        f"Formula referenced {sorted(deps['external'])}; re-derived "
                        "from the local sheet of the same name.", sheet, cell)
        if isinstance(recalculated, ExcelError):
            out["status"] = "invalid"
            reasons.append(f"formula_error:{recalculated.code}")
            self._issue("error", "formula_error",
                        f"Formula re-derives to {recalculated.code}.", sheet, cell)
            return self._finish(out, cfg_metric)

        if isinstance(recalculated, str):
            # e.g. the workbook's own "n/a" / "n/m (...)" guard results.
            if _is_unavailable(recalculated, cfg):
                out["status"] = "unavailable"
                reasons.append("not_provided_by_workbook")
            else:
                out["status"] = "invalid"
                reasons.append("inappropriate_value:text_result")
            return self._finish(out, cfg_metric)

        value = float(recalculated)
        if cached_is_error:
            reasons.append(f"cached_error_repaired:{cached}")
            self._corrections.append(
                {
                    "kind": "cached_error_recomputed",
                    "reference": f"{sheet}!{cell}",
                    "cached_value": cached,
                    "recalculated_value": value,
                    "rationale": "Cached result was an Excel error; the value was "
                                 "re-derived from documented local inputs.",
                }
            )
        elif isinstance(cached, (int, float)) and not isinstance(cached, bool):
            matched = _same(cached, value, tol["relative_tolerance"], tol["absolute_tolerance"])
            out["reconciliation"] = {
                "cached_value": float(cached),
                "recalculated_value": value,
                "absolute_difference": abs(float(cached) - value),
                "matched": matched,
            }
            if not matched:
                out["status"] = "invalid"
                reasons.append("recalculation_mismatch")
                self._issue(
                    "error", "recalculation_mismatch",
                    f"Cached {cached!r} differs from independently re-derived {value!r}.",
                    sheet, cell,
                )
                return self._finish(out, cfg_metric)
        else:
            reasons.append("no_numeric_cached_value_to_reconcile")
            self._issue("warning", "no_cached_value",
                        "No numeric cached result to reconcile against; using the "
                        "re-derived value.", sheet, cell)

        out["value"] = value
        placeholder = self._placeholder_assumptions(deps["assumptions"])
        if placeholder:
            out["status"] = "assumed"
            reasons.extend(f"depends_on_placeholder_assumption:{k}" for k in placeholder)
        else:
            out["status"] = "derived"
            reasons.extend(
                f"depends_on_methodology_switch:{k}"
                for k in self._assumption_keys(deps["assumptions"], "methodology_switch")
            )
        return self._finish(out, cfg_metric)

    def _placeholder_assumptions(self, keys: set[str]) -> list[str]:
        return self._assumption_keys(keys, "placeholder")

    def _assumption_keys(self, keys: set[str], kind: str) -> list[str]:
        kinds = {a["key"]: a["kind"] for a in self._assumptions}
        return sorted(k for k in keys if kinds.get(k) == kind)

    def _finish(self, out: dict[str, Any], cfg_metric: dict[str, Any]) -> dict[str, Any]:
        """Plausibility + eligibility for the declared use."""
        value, reasons = out["value"], out["reasons"]
        bounds = cfg_metric.get("plausible_range")
        if value is not None and bounds and out["status"] in {"source_aggregate", "derived", "assumed"}:
            low, high = bounds
            if not (low <= value <= high):
                out["status"] = "invalid"
                reasons.append(f"inappropriate_value:outside_plausible_range[{low},{high}]")
                self._issue("error", "inappropriate_value",
                            f"Value {value!r} is outside the configured plausible "
                            f"range [{low}, {high}].", out["sheet"], out["cell"])
        if out["status"] == "invalid":
            out["value"] = None
        use = cfg_metric.get("use")
        eligible = out["status"] in {"source_aggregate", "derived"} and use == "comparison"
        out["eligible_for_comparison"] = eligible
        if not eligible:
            if use is None:
                out["exclusion_reason"] = "row_not_mapped_to_a_benchmark_metric"
            elif use != "comparison":
                out["exclusion_reason"] = f"declared_use:{use}"
            elif out["status"] == "assumed":
                out["exclusion_reason"] = "depends_on_placeholder_assumption"
            else:
                out["exclusion_reason"] = f"status:{out['status']}"
        else:
            out["exclusion_reason"] = None
        return out

    def _check_row_consistency(self, entry, forms: dict[str, str], sheet: str) -> None:
        if len(set(forms.values())) <= 1:
            return
        counts: dict[str, int] = {}
        for form in forms.values():
            counts[form] = counts.get(form, 0) + 1
        majority = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))[0][0]
        for industry, form in forms.items():
            if form == majority:
                continue
            value = entry["values"][industry]
            value["status"] = "invalid"
            value["value"] = None
            value["eligible_for_comparison"] = False
            value["exclusion_reason"] = "status:invalid"
            value["reasons"].append("inconsistent_formula")
            self._issue(
                "error", "inconsistent_formula",
                f"Formula differs from the other industry columns in row "
                f"{entry['row']} ({entry['label']!r}).", sheet, value["cell"],
            )

    def _read_sic_codes(self) -> dict[str, Any]:
        cfg = self._mapping["sic_sheet"]
        sheet = self._mapping["sheets"]["sic_codes"]
        ws = self._formulas[sheet]
        columns = self._header_columns(ws, cfg["header_row"], cfg["columns"], sheet)
        records = []
        by_code: dict[int, int] = {}
        titles: dict[str, list[int]] = {}
        for row in range(cfg["header_row"] + 1, ws.max_row + 1):
            code = ws[f"{columns['sic_code']}{row}"].value
            if code is None:
                continue
            record = {"row": row, "status": "observed", "sheet": sheet,
                      "cell": f"{columns['sic_code']}{row}"}
            for key, column in columns.items():
                record[key] = ws[f"{column}{row}"].value
            if not isinstance(code, int) or not (100 <= code <= 9999):
                record["status"] = "invalid"
                self._issue("error", "inappropriate_value",
                            f"SIC code {code!r} is not a 3-4 digit integer.", sheet, record["cell"])
            for required in ("segment", "title", "mapping_confidence"):
                if record.get(required) in (None, ""):
                    record["status"] = "invalid"
                    self._issue("error", "blank_mandatory_cell",
                                f"SIC row is missing {required!r}.", sheet,
                                f"{columns[required]}{row}")
            if isinstance(code, int):
                if code in by_code:
                    self._issue("error", "duplicate_sic_code",
                                f"SIC code {code} repeats (rows {by_code[code]} and {row}).",
                                sheet, record["cell"])
                by_code[code] = row
                titles.setdefault(str(record.get("title")), []).append(code)
            records.append(record)
        for title, codes in sorted(titles.items()):
            if len(codes) > 1:
                self._issue(
                    "warning", "duplicate_industry_title",
                    f"Industry title {title!r} is recorded for several SIC codes "
                    f"{sorted(codes)}; at least one title is likely mis-transcribed.", sheet,
                )
        notes = [
            {"sheet": sheet, "cell": f"A{r}", "text": ws[f"A{r}"].value}
            for r in cfg["documentation_rows"] if ws[f"A{r}"].value
        ]
        return {"records": records, "documentation": notes}

    def _read_companies(self, industries: list[str]) -> dict[str, Any]:
        cfg = self._mapping["companies_sheet"]
        sheet = self._mapping["sheets"]["companies"]
        ws = self._formulas[sheet]
        columns = self._header_columns(ws, cfg["header_row"], cfg["columns"], sheet)
        aliases = self._mapping.get("industry_aliases", {})
        records, by_industry = [], {}
        alias_hits: dict[str, int] = {}
        letters = {key: column_index_from_string(col) - 1 for key, col in columns.items()}
        for index, row in enumerate(
            ws.iter_rows(min_row=cfg["header_row"] + 1, values_only=True),
            start=cfg["header_row"] + 1,
        ):
            name = row[letters["name"]] if letters["name"] < len(row) else None
            if name is None:
                continue
            record = {key: (row[i] if i < len(row) else None) for key, i in letters.items()}
            original = record["industry_group"]
            if original in aliases:
                alias_hits[original] = alias_hits.get(original, 0) + 1
                record["industry_group"] = aliases[original]
            record.update({"row": index, "record_type": "company_identity",
                           "status": "observed", "has_financial_observations": False})
            if record["industry_group"] not in industries:
                record["status"] = "invalid"
                self._issue("error", "inappropriate_value",
                            f"Company industry group {original!r} is not one of the "
                            "six benchmark industries.", sheet, f"{columns['industry_group']}{index}")
            by_industry[record["industry_group"]] = by_industry.get(record["industry_group"], 0) + 1
            records.append(record)
        for original, count in sorted(alias_hits.items()):
            self._issue(
                "info", "industry_label_normalized",
                f"{count} company rows use the label {original!r}; normalized to "
                f"{aliases[original]!r} via the mapping alias.", sheet,
            )
        return {
            "record_type": "company_identity",
            "note": (
                "Identity and classification only (name, ticker, industry group, "
                "SIC). The sheet holds NO company financial observations and cannot "
                "populate an empirical peer cohort."
            ),
            "count": len(records),
            "by_industry": dict(sorted(by_industry.items())),
            "records": records,
        }

    def _header_columns(self, ws, header_row, wanted, sheet) -> dict[str, str]:
        found = {str(c.value): c.column_letter for c in ws[header_row] if c.value is not None}
        columns = {}
        for key, header in wanted.items():
            if header not in found:
                raise WorkbookImportError(
                    f"Sheet {sheet!r} is missing the required column {header!r}."
                )
            columns[key] = found[header]
        return columns

    def _read_documentation(self) -> dict[str, Any]:
        cfg = self._mapping["assumptions_sheet"]
        sheet = self._mapping["sheets"]["assumptions"]
        ws = self._formulas[sheet]
        first, last = cfg["limitations_rows"]
        limitations = [
            {"sheet": sheet, "cell": f"B{r}", "title": ws[f"A{r}"].value, "text": ws[f"B{r}"].value}
            for r in range(first, last + 1) if ws[f"A{r}"].value
        ]
        first, last = cfg["sources_rows"]
        sources = []
        for r in range(first, last + 1):
            label, url = ws[f"A{r}"].value, ws[f"B{r}"].value
            if label is None and url is None:
                continue
            valid = isinstance(url, str) and _URL.fullmatch(url.strip()) is not None
            if not valid:
                self._issue("error", "missing_source_documentation",
                            f"Source {label!r} has no URL.", sheet, f"B{r}")
            sources.append({"sheet": sheet, "cell": f"B{r}", "label": label, "url": url})
        if not sources:
            self._issue("error", "missing_source_documentation",
                        "No source documentation rows found.", sheet)
        if not limitations:
            self._issue("error", "missing_source_documentation",
                        "No assumptions/limitations text found.", sheet)
        return {"limitations": limitations, "sources": sources}

    def _read_vintage(self) -> dict[str, Any]:
        statements = []
        found: set[str] = set()
        for item in self._mapping["vintage_cells"]:
            text = self._formulas[item["sheet"]][item["cell"]].value
            match = _VINTAGE.search(text or "")
            statements.append({**item, "text": text, "parsed": _month(match) if match else None})
            if match:
                found.add(_month(match))
            else:
                self._issue("error", "missing_source_date",
                            "No 'as of <Month YYYY>' data date found.", item["sheet"], item["cell"])
        if len(found) > 1:
            self._issue("error", "inconsistent_source_dates",
                        f"Workbook states several data dates: {sorted(found)}.")
        sic_sheet = self._mapping["sheets"]["sic_codes"]
        sic_updated = None
        for row in self._mapping["sic_sheet"]["documentation_rows"]:
            match = _UPDATED.search(self._formulas[sic_sheet][f"A{row}"].value or "")
            if match:
                sic_updated = (
                    f"{match.group(3)}-{_month_number(match.group(1)):02d}-"
                    f"{int(match.group(2)):02d}"
                )
        if sic_updated is None:
            self._issue("warning", "missing_source_date",
                        "SIC code list has no 'last updated' date.", sic_sheet)
        return {
            "data_as_of": sorted(found)[0] if len(found) == 1 else None,
            "statements": statements,
            "sic_list_last_updated": sic_updated,
        }

    def _cross_checks(self, metrics: list[dict[str, Any]]) -> None:
        by_label = {m["label"]: m for m in metrics}
        for check in self._mapping.get("cross_checks", []):
            metric = by_label.get(check["label"])
            if metric is None:
                self._issue("warning", "cross_check_missing",
                            f"Cross-check row {check['label']!r} not found.")
                continue
            for industry, value in metric["values"].items():
                number = value["value"]
                if number is None or abs(number) > check["max_abs"]:
                    self._issue(
                        "warning", "cross_check_failed",
                        f"{check['label']} for {industry} is {number!r}, outside "
                        f"±{check['max_abs']}; derived sales (and every ratio built "
                        "on them) are less reliable for this industry.",
                        value["sheet"], value["cell"],
                    )

    def _check_sample_sizes(self, metrics, companies) -> None:
        sample = next((m for m in metrics if m["label"].startswith("Firms in sample")), None)
        if sample is None:
            return
        for industry, value in sample["values"].items():
            listed = companies["by_industry"].get(industry, 0)
            if value["value"] is not None and int(value["value"]) != listed:
                self._issue(
                    "info", "company_list_differs_from_sample",
                    f"{industry}: the aggregate covers {int(value['value'])} firms but "
                    f"the company sheet lists {listed}; the list is not the "
                    "aggregate's sample.",
                    value["sheet"], value["cell"],
                )

    # --------------------------------------------------------------- report

    def _issue(self, severity, code, message, sheet=None, cell=None) -> None:
        issue = Issue(severity, code, message, sheet, cell)
        if issue.as_dict() not in [i.as_dict() for i in self._issues]:
            self._issues.append(issue)

    def _report(self, dataset: dict[str, Any]) -> dict[str, Any]:
        status_counts = {s: 0 for s in VALUE_STATUSES}
        eligible, excluded = [], []
        for metric in dataset["metrics"]:
            for industry, value in metric["values"].items():
                status_counts[value["status"]] += 1
                row = {
                    "benchmark_metric_id": metric["benchmark_metric_id"],
                    "label": metric["label"],
                    "section": metric["section"],
                    "industry": industry,
                    "sheet": value["sheet"],
                    "cell": value["cell"],
                    "status": value["status"],
                    "declared_use": metric["declared_use"],
                    "value": value["value"],
                }
                if value["eligible_for_comparison"]:
                    eligible.append(row)
                else:
                    excluded.append({**row, "exclusion_reason": value["exclusion_reason"],
                                     "reasons": value["reasons"]})
        status_counts["observed"] += sum(
            1 for r in dataset["sic_codes"]["records"] if r["status"] == "observed"
        ) + sum(1 for r in dataset["companies"]["records"] if r["status"] == "observed")
        status_counts["assumed"] += len(dataset["assumptions"])
        issues = [i.as_dict() for i in self._issues]
        severities = {s: sum(1 for i in issues if i["severity"] == s)
                      for s in ("error", "warning", "info")}
        return {
            "schema_version": "workbook-validation-report-1.0",
            "importer_version": IMPORTER_VERSION,
            "source_sha256": dataset["source"]["sha256"],
            "dataset_hash": dataset["dataset_hash"],
            "mapping_hash": dataset["mapping_hash"],
            "sheets_imported": dataset["source"]["sheets"],
            "industries": dataset["industries"],
            "vintage": dataset["vintage"],
            "status_counts": status_counts,
            "issue_counts": severities,
            "issues": issues,
            "eligible_values": eligible,
            "excluded_values": excluded,
            "correction_log": dataset["correction_log"],
            "acceptance": {
                "every_value_classified": all(
                    v["status"] in VALUE_STATUSES
                    and (v["eligible_for_comparison"] or v["exclusion_reason"])
                    for m in dataset["metrics"] for v in m["values"].values()
                ),
                "eligible_count": len(eligible),
                "excluded_count": len(excluded),
            },
        }


def _is_unavailable(text: str, cfg: dict[str, Any]) -> bool:
    lowered = text.strip().lower()
    return lowered in cfg["unavailable_literals"] or any(
        lowered.startswith(p) for p in cfg["unavailable_prefixes"]
    )


def _same(a: Any, b: Any, rel: float, absolute: float) -> bool:
    try:
        x, y = float(a), float(b)
    except (TypeError, ValueError):
        return a == b
    return abs(x - y) <= max(absolute, rel * max(abs(x), abs(y)))


def _jsonable(value: Any) -> Any:
    if isinstance(value, ExcelError):
        return value.code
    if isinstance(value, (int, float, str, bool)) or value is None:
        return value
    return str(value)


def _month_number(name: str) -> int:
    return _MONTHS.split("|").index(name.lower()) + 1


def _month(match: re.Match[str]) -> str:
    return f"{match.group(2)}-{_month_number(match.group(1)):02d}"


def import_workbook(
    data: bytes,
    sector_config: dict[str, Any],
    *,
    filename: str = "workbook.xlsx",
) -> WorkbookImport:
    """Import ``data`` under a sector configuration's workbook mapping."""
    return HealthcareWorkbookImporter(
        sector_config["workbook_mapping"],
        benchmark_metrics=sector_config.get("benchmark_metrics", {}),
    ).import_bytes(data, filename=filename)


def industry_values(dataset: dict[str, Any], industry: str) -> dict[str, dict[str, Any]]:
    """``{benchmark_metric_id: value-entry}`` for one industry column.

    Each entry is the value record plus its row-level label/definition/note.
    Rows not mapped to a benchmark metric id are omitted.
    """
    if industry not in dataset["industries"]:
        raise KeyError(f"Industry {industry!r} is not in the reference dataset.")
    out: dict[str, dict[str, Any]] = {}
    for metric in dataset["metrics"]:
        metric_id = metric["benchmark_metric_id"]
        if metric_id is None:
            continue
        out[metric_id] = {
            **metric["values"][industry],
            "label": metric["label"],
            "definition": metric["definition"],
            "note": metric["note"],
            "section": metric["section"],
            "declared_use": metric["declared_use"],
            "units": metric["units"],
        }
    return out
